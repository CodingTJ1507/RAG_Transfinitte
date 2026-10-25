import csv
import hashlib
import json
import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
import pymupdf
from docx import Document

from app.config import settings
from app.services.gemini import describe_image

logger = logging.getLogger(__name__)
IMAGE_MIME_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}

@dataclass
class ParsedSource:
    source: dict[str, Any]
    source_hash: str
    units: list[dict[str, Any]]


def load_manifest(dataset_dir: Path | None = None) -> tuple[Path, dict[str, Any]]:
    root = (dataset_dir or settings.dataset_path).resolve()
    manifest_path = root / "metadata" / "source_manifest.json"
    if not manifest_path.is_file():
        raise FileNotFoundError(f"Dataset manifest not found: {manifest_path}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if not isinstance(manifest.get("sources"), list):
        raise ValueError("Dataset manifest must contain a sources list")
    return root, manifest


def _safe_source_path(root: Path, relative_path: str) -> Path:
    path = (root / relative_path).resolve()
    try:
        path.relative_to(root.resolve())
    except ValueError as exc:
        raise ValueError(f"Manifest path escapes the dataset directory: {relative_path}") from exc
    if not path.is_file():
        raise FileNotFoundError(f"Manifest source does not exist: {relative_path}")
    if path.stat().st_size > settings.MAX_SOURCE_FILE_BYTES:
        raise ValueError(f"Source exceeds MAX_SOURCE_FILE_BYTES: {relative_path}")
    return path


def _source_hash(path: Path, declared_hash: str | None) -> str:
    raw_content = path.read_bytes()
    raw_hash = hashlib.sha256(raw_content).hexdigest()
    if path.suffix.lower() in {".html", ".htm", ".vtt", ".txt", ".md", ".json", ".csv"}:
        normalized_hash = hashlib.sha256(raw_content.replace(b"\r\n", b"\n")).hexdigest()
    else:
        normalized_hash = raw_hash
    if declared_hash:
        if declared_hash not in {raw_hash, normalized_hash}:
            raise ValueError(f"Source content hash does not match the manifest: {path.name}")
        return declared_hash
    return normalized_hash


def _text_unit(text: str, locator: dict[str, Any], **metadata: Any) -> dict[str, Any] | None:
    cleaned = re.sub(r"\s+", " ", text).strip()
    if not cleaned:
        return None
    return {"text": cleaned, "locator": locator, "metadata": metadata}


PDF_SECTION_SIZE = 16
PDF_SUBSECTION_SIZE = 12
PDF_RUNNING_TEXT_MAX_SIZE = 8
PDF_HEADER_ZONE = 0.06  # running header and footer sit above / below these fractions of the page height
PDF_FOOTER_ZONE = 0.93
PDF_DOT_LEADER = re.compile(r"(?:\.\s+){3,}")


def _pdf_lines(page: Any) -> list[dict[str, Any]]:
    lines: list[dict[str, Any]] = []
    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = [span for span in line.get("spans", []) if span["text"].strip()]
            text = "".join(span["text"] for span in spans).strip()
            if not text:
                continue
            bold = all("Bold" in span["font"] for span in spans)
            sizes = {round(span["size"]) for span in spans}
            level = 1 if bold and sizes == {PDF_SECTION_SIZE} else 2 if bold and sizes == {PDF_SUBSECTION_SIZE} else 0
            lines.append({"text": text, "top": line["bbox"][1], "bold": bold, "max_size": max(sizes), "level": level})
    return lines


def _pdf_unit(page_number: int, section_path: list[str], texts: list[str]) -> dict[str, Any] | None:
    return _text_unit(" ".join(texts), {
        "page": page_number,
        "section": section_path[-1] if section_path else "",
        "section_path": list(section_path),
    })


def _cover_unit(lines: list[dict[str, Any]]) -> dict[str, Any] | None:
    owner = next(
        (lines[index + 1]["text"] for index, line in enumerate(lines[:-1]) if line["text"] == "Owner" and line["bold"]),
        None,
    )
    return _text_unit(
        " ".join(line["text"] for line in lines),
        {"page": 1, "section": "Cover", "section_path": ["Cover"]},
        **({"cover_owner": owner} if owner else {}),
    )


def _is_contents_page(lines: list[dict[str, Any]]) -> bool:
    dot_leaders = sum(bool(PDF_DOT_LEADER.search(line["text"])) for line in lines)
    return any(line["text"] == "Contents" for line in lines) and dot_leaders >= 3


def _without_running_text(lines: list[dict[str, Any]], page_height: float) -> list[dict[str, Any]]:
    return [
        line for line in lines
        if not (
            line["max_size"] <= PDF_RUNNING_TEXT_MAX_SIZE
            and (line["top"] < page_height * PDF_HEADER_ZONE or line["top"] > page_height * PDF_FOOTER_ZONE)
        )
    ]


def _section_units(
    lines: list[dict[str, Any]], page_number: int, section_path: list[str],
) -> tuple[list[dict[str, Any]], list[str]]:
    """One unit per section_path on this page; also returns the path that carries to the next page."""
    groups: list[tuple[list[str], list[str], bool]] = []  # (section_path, texts, started by a heading)
    for line in lines:
        if line["level"]:
            heading = line["text"]
            section_path = [heading] if line["level"] == 1 or not section_path else [section_path[0], heading]
            groups.append((list(section_path), [heading], True))
        elif groups:
            groups[-1][1].append(line["text"])
        else:
            groups.append((list(section_path), [line["text"]], False))
    units = [
        _pdf_unit(page_number, group_path, texts)
        for group_path, texts, from_heading in groups
        if not (from_heading and len(texts) == 1)  # a heading directly followed by another heading
    ]
    return [unit for unit in units if unit], section_path


def _parse_pdf(path: Path) -> list[dict[str, Any]]:
    with pymupdf.open(path) as document:
        pages = [(page.rect.height, _pdf_lines(page)) for page in document]
    has_headings = any(line["level"] for _, lines in pages for line in lines)
    units: list[dict[str, Any]] = []
    section_path: list[str] = []
    for page_number, (height, lines) in enumerate(pages, start=1):
        if page_number == 1:
            page_units = [_cover_unit(lines)]
        elif _is_contents_page(lines):
            continue
        elif not has_headings:
            page_units = [_pdf_unit(page_number, [], [line["text"] for line in _without_running_text(lines, height)])]
        else:
            page_units, section_path = _section_units(_without_running_text(lines, height), page_number, section_path)
        units.extend(unit for unit in page_units if unit)
    return units


def _parse_docx(path: Path) -> list[dict[str, Any]]:
    document = Document(path)
    units: list[dict[str, Any]] = []
    section = "Document"
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name.startswith("Heading"):
            section = text
        unit = _text_unit(
            text,
            {"section": section, "paragraph": paragraph_number},
        )
        if unit:
            units.append(unit)
    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            values = [cell.text.strip() for cell in row.cells]
            unit = _text_unit(
                " | ".join(value for value in values if value),
                {"section": f"Table {table_number}", "row": row_number},
            )
            if unit:
                units.append(unit)
    return units


def _parse_markdown(path: Path) -> list[dict[str, Any]]:
    units: list[dict[str, Any]] = []
    section = "Document"
    for line_number, line in enumerate(path.read_text(encoding="utf-8", errors="replace").splitlines(), 1):
        clean = line.strip()
        if not clean:
            continue
        heading = re.match(r"^#{1,6}\s+(.+?)\s*#*\s*$", clean)
        if heading:
            section = heading.group(1).strip()
            content = section
        else:
            content = clean
        unit = _text_unit(content, {"section": section, "line": line_number})
        if unit:
            units.append(unit)
    return units


def _html_text(element: Any) -> str:
    if element.name == "table":
        return " | ".join(
            cell.get_text(" ", strip=True)
            for cell in element.find_all(["th", "td"])
        )
    return element.get_text(" ", strip=True)


def _slug(text: str) -> str:
    return re.sub(r"\s+", "-", re.sub(r"[^a-z0-9 ]", "", text.lower()))


def _parse_html(path: Path) -> list[dict[str, Any]]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for element in soup(["script", "style", "nav", "footer", "header", "noscript", "svg"]):
        element.decompose()
    main = soup.find("main") or soup.find("article") or soup.body or soup
    # A block inside another block (a p in an li) is already part of its parent's text.
    elements = [
        element for element in main.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "blockquote", "table"])
        if not element.find_parent(["p", "li", "pre", "blockquote", "table"])
    ]
    page_title = soup.title.get_text(" ", strip=True) if soup.title else ""
    h1 = main.find("h1")
    intro_section = h1.get_text(" ", strip=True) if h1 else page_title

    h2_elements = [element for element in elements if element.name == "h2"]
    if not h2_elements:
        unit = _text_unit(
            " ".join(_html_text(element) for element in elements),
            {"section": intro_section or "Webpage content", "html_id": "intro"},
        )
        return [unit] if unit else []

    groups: list[tuple[str, str, list[Any]]] = []
    current_heading = ""
    current_id = ""
    current_elements: list[Any] = []

    for element in elements:
        if element.name == "h2":
            if current_elements:
                groups.append((current_heading, current_id, current_elements))
            current_heading = _html_text(element)
            current_id = element.get("id") or _slug(current_heading)
            current_elements = [element]
        elif current_elements:
            current_elements.append(element)

    if current_elements:
        groups.append((current_heading, current_id, current_elements))

    units: list[dict[str, Any]] = []
    for heading, html_id, group_elements in groups:
        unit = _text_unit(
            " ".join(_html_text(element) for element in group_elements),
            {"section": heading or "Webpage content", "html_id": html_id},
        )
        if unit:
            units.append(unit)

    first_h2_index = next(
        index for index, element in enumerate(elements)
        if element is h2_elements[0]
    )
    intro_elements = elements[:first_h2_index]
    intro_texts = [_html_text(element) for element in intro_elements]
    intro_content = [
        text for element, text in zip(intro_elements, intro_texts)
        if element is not h1 and text
    ]
    if intro_content:
        intro_unit = _text_unit(
            " ".join(intro_texts),
            {"section": intro_section or "Webpage content", "html_id": "intro"},
        )
        if intro_unit:
            units.insert(0, intro_unit)

    return units


def _parse_vtt(path: Path) -> list[dict[str, Any]]:
    content = path.read_text(encoding="utf-8", errors="replace").replace("\r\n", "\n")
    blocks = re.split(r"\n\s*\n", content)
    units: list[dict[str, Any]] = []
    for block in blocks:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        timing_index = next((i for i, line in enumerate(lines) if "-->" in line), None)
        if timing_index is None:
            continue
        cue_id = lines[timing_index - 1] if timing_index > 0 else None
        timing = lines[timing_index].split("-->", 1)
        start = timing[0].strip().split(" ", 1)[0]
        end = timing[1].strip().split(" ", 1)[0]
        body = " ".join(lines[timing_index + 1:])
        # WebVTT voice tag "<v Name>text" first, then the plain "Name: text" style.
        speaker_match = (
            re.match(r"^<v(?:\.[\w.-]+)?\s+([^>]+)>(.*?)(?:</v>)?$", body)
            or re.match(r"^([A-Za-z][A-Za-z0-9 _.-]{0,40}):\s*(.*)$", body)
        )
        speaker = speaker_match.group(1).strip() if speaker_match else None
        if speaker_match:
            body = speaker_match.group(2)
        unit = _text_unit(body, {"cue_id": cue_id, "start": start, "end": end}, speaker=speaker)
        if unit:
            units.append(unit)
    return units


TXT_HEADING_MAX_LENGTH = 40
TXT_PART_MAX_CHARS = 2000
TXT_ITEM = re.compile(r"^(-|\*|\d+\.)\s")


def _is_text_heading(line: str) -> bool:
    clean = line.strip()
    return (
        len(clean) < TXT_HEADING_MAX_LENGTH
        and re.search(r"[A-Za-z]", clean) is not None
        and clean.upper() == clean
    )


def _text_paragraphs(lines: list[str]) -> list[tuple[int, str]]:
    paragraphs: list[tuple[int, str]] = []
    current: list[str] = []
    start = 0
    for line_number, line in enumerate(lines, start=1):
        clean = line.strip()
        if clean:
            if not current:
                start = line_number
            current.append(clean)
        elif current:
            paragraphs.append((start, " ".join(current)))
            current = []
    if current:
        paragraphs.append((start, " ".join(current)))
    return paragraphs


def _pack_parts(paragraphs: list[tuple[int, str]], limit: int = TXT_PART_MAX_CHARS) -> list[dict[str, Any]]:
    parts: list[tuple[int, str]] = []
    for line_number, paragraph in paragraphs:
        if parts and len(parts[-1][1]) + 1 + len(paragraph) <= limit:
            parts[-1] = (parts[-1][0], f"{parts[-1][1]} {paragraph}")
        else:
            parts.append((line_number, paragraph))
    units = [
        _text_unit(text, {"section": f"Part {number}", "line": start})
        for number, (start, text) in enumerate(parts, start=1)
    ]
    return [unit for unit in units if unit]


def _section_items(lines: list[tuple[int, str]]) -> list[tuple[int, list[str]]]:
    """A section's bullet or numbered items with their wrapped lines; plain prose stays one item."""
    items: list[tuple[int, list[str]]] = []
    for line_number, text in lines:
        if TXT_ITEM.match(text) or not items:
            items.append((line_number, [text]))
        else:
            items[-1][1].append(text)
    return items


def _text_sections(content_lines: list[tuple[int, str]]) -> list[dict[str, Any]]:
    """One unit per item, prefixed with its section, so one note's unrelated points stay searchable apart."""
    sections: list[tuple[str, list[tuple[int, str]]]] = []
    for line_number, line in content_lines:
        clean = line.strip()
        if not clean:
            continue
        if _is_text_heading(clean):
            sections.append((clean, []))
            continue
        if not sections:
            sections.append(("HEADER", []))
        sections[-1][1].append((line_number, clean))
    units = [
        _text_unit(f"{name}: {' '.join(texts)}", {"section": name, "line": start})
        for name, lines in sections
        for start, texts in _section_items(lines)
    ]
    return [unit for unit in units if unit]


def _parse_text(path: Path) -> list[dict[str, Any]]:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    title_index = next((index for index, line in enumerate(lines) if line.strip()), None)
    if title_index is None:
        return []
    content_lines = list(enumerate(lines[title_index + 1:], title_index + 2))
    if any(_is_text_heading(line) for _, line in content_lines):
        return _text_sections(content_lines)
    return _pack_parts(_text_paragraphs(lines))


def _structured_records(value: Any) -> list[tuple[str, dict[str, Any]]]:
    records: list[tuple[str, dict[str, Any]]] = []
    if isinstance(value, list):
        candidates = [("record", item) for item in value if isinstance(item, dict)]
    elif isinstance(value, dict):
        candidates = []
        for key, child in value.items():
            if isinstance(child, list):
                candidates.extend((key, item) for item in child if isinstance(item, dict))
        if not candidates:
            candidates = [("record", value)]
    else:
        candidates = [("record", {"value": value})]
    for collection_name, item in candidates:
        identifier_key = next(
            (key for key in ("log_id", "event_id", "record_id", "dependency_id", "service_id", "id")
             if item.get(key) is not None),
            None,
        )
        identifier = str(item[identifier_key]) if identifier_key else f"{collection_name}-{len(records) + 1}"
        records.append((identifier, item))
    return records


def _parse_json(path: Path) -> list[dict[str, Any]]:
    value = json.loads(path.read_text(encoding="utf-8"))
    units: list[dict[str, Any]] = []
    for record_id, record in _structured_records(value):
        text = " | ".join(
            f"{key}: {json.dumps(item, ensure_ascii=False) if isinstance(item, (dict, list)) else item}"
            for key, item in record.items()
        )
        timestamp = next(
            (record[key] for key in ("timestamp", "created_at", "received_at") if record.get(key)),
            None,
        )
        locator_key = next(
            (key for key in ("log_id", "event_id", "record_id", "dependency_id", "service_id", "id")
             if record.get(key) is not None),
            "record_id",
        )
        locator_value = record.get(locator_key, record_id)
        unit = _text_unit(
            text,
            {locator_key: locator_value, "timestamp": timestamp},
            record=record,
        )
        if unit:
            units.append(unit)
    return units


def _parse_csv(path: Path) -> list[dict[str, Any]]:
    units: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row_number, record in enumerate(csv.DictReader(handle), start=2):
            text = " | ".join(f"{key}: {value}" for key, value in record.items() if value is not None)
            row_id = record.get("row_id") or row_number
            timestamp = record.get("timestamp")
            unit = _text_unit(
                text,
                {"row_id": row_id, "line": row_number, "timestamp": timestamp},
                record=record,
            )
            if unit:
                units.append(unit)
    return units


def _source_owner(source: dict[str, Any], suffix: str, units: list[dict[str, Any]]) -> str:
    """First of: the manifest or upload owner, the PDF cover owner, the first service's team."""
    owner = source.get("owner")
    if isinstance(owner, str) and owner.strip():
        return owner
    cover_owner = units[0].get("metadata", {}).get("cover_owner") if suffix == ".pdf" and units else None
    if isinstance(cover_owner, str) and cover_owner.strip():
        return cover_owner
    services = source.get("services") or []
    if isinstance(services, list) and services and isinstance(services[0], str):
        return settings.service_owner_map.get(services[0], "Unassigned")
    return "Unassigned"


def _source_authority(source: dict[str, Any]) -> str:
    """How far the evidence can be trusted, from the manifest tags and source type (first rule wins)."""
    tags = set(source.get("tags") or [])
    source_type = source.get("source_type")
    if tags & {"hypothesis", "initial-report"}:
        return "early_hypothesis"
    if "background" in tags:
        return "background"
    if source_type in {"structured", "structured_record"}:
        return "observed"
    if source_type == "image":
        return "reference" if "architecture" in tags else "observed"
    if source_type == "transcript":
        return "discussion"
    return "reference"


def _read_image(path: Path, suffix: str) -> dict[str, Any] | None:
    if not settings.IMAGE_ANALYSIS_ENABLED:
        return None
    try:
        return describe_image(path.read_bytes(), IMAGE_MIME_TYPES[suffix])
    except RuntimeError:
        logger.warning("Gemini image reading failed for %s; falling back to metadata", path.name)
        return None


def _parse_image(path: Path, source: dict[str, Any], suffix: str) -> list[dict[str, Any]]:
    from PIL import Image

    with Image.open(path) as image:
        image.verify()
    title = source.get("title", path.name)
    description = source.get("description", "")
    reading = _read_image(path, suffix)
    if reading is None:
        unit = _text_unit(
            f"{title}. Image metadata description: {description}",
            {"image": path.name, "analysis": "metadata_only", "ocr_analyzed": False},
        )
    else:
        manifest_part = f" {description}." if description else ""
        unit = _text_unit(
            f"{title}.{manifest_part} Description: {reading['description']} "
            f"Visible text: {' | '.join(reading['visible_text'])}",
            {"image": path.name, "analysis": "gemini_vision"},
            visible_text=reading["visible_text"],
        )
    return [unit] if unit else []


def parse_source(root: Path, source: dict[str, Any]) -> ParsedSource:
    relative_path = source.get("file_path")
    if not isinstance(relative_path, str):
        raise ValueError("Manifest source is missing a file_path")
    path = _safe_source_path(root, relative_path)
    source_hash = _source_hash(path, source.get("sha256"))
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        units = _parse_pdf(path)
    elif suffix in {".html", ".htm"}:
        units = _parse_html(path)
    elif suffix == ".vtt":
        units = _parse_vtt(path)
    elif suffix == ".md":
        units = _parse_markdown(path)
    elif suffix == ".docx":
        units = _parse_docx(path)
    elif suffix == ".txt":
        units = _parse_text(path)
    elif suffix == ".json":
        units = _parse_json(path)
    elif suffix == ".csv":
        units = _parse_csv(path)
    elif suffix in IMAGE_MIME_TYPES:
        units = _parse_image(path, source, suffix)
    else:
        raise ValueError(f"Unsupported source format: {suffix or 'unknown'}")

    source_record = {
        **source,
        "file_path": relative_path.replace("\\", "/"),
        "source_hash": source_hash,
        "size_bytes": path.stat().st_size,
        "owner": _source_owner(source, suffix, units),
        "authority": _source_authority(source),
        "analysis_status": (
            (
                "vision_analyzed"
                if units and units[0]["locator"].get("analysis") == "gemini_vision"
                else "metadata_only"
            )
            if suffix in IMAGE_MIME_TYPES
            else ("parsed" if units else "no_extractable_text")
        ),
    }
    return ParsedSource(source=source_record, source_hash=source_hash, units=units)
