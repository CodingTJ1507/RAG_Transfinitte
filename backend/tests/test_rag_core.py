import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from fastapi.testclient import TestClient
from docx import Document

from app.api.routes_eval import retrieval_metrics
from app.ingestion.chunking import chunk_document_with_location, chunk_evidence_units
from app.ingestion.parsers import load_manifest, parse_source
from app.models.schemas import Citation, SourceType, SupportedClaim
from app.reasoning.investigator import investigate_query, verify_claims
from app.reasoning.query_transform import decompose_query, rewrite_query
from app.retrieval.hybrid_search import _atlas_search_filters, reciprocal_rank_fusion
from app.retrieval.vector_search import build_vector_filter
from app.services.embeddings import embed_query, embed_texts
from app.services.gemini import generate_json
from app.main import app
from app.config import Settings
from app.core.database import SearchIndexNotReadyError


class DatasetContractTests(unittest.TestCase):
    def test_ground_truth_points_to_canonical_dataset_incident(self):
        backend_root = Path(__file__).resolve().parents[1]
        legacy = json.loads((Path(__file__).parent / "ground_truth.json").read_text(encoding="utf-8"))
        self.assertEqual(legacy["incident_id"], "INC-2026-1042")
        root, manifest = load_manifest(backend_root / "synthetic_data")
        self.assertEqual(manifest["incident_id"], legacy["incident_id"])
        known_sources = {entry["source_id"] for entry in manifest["sources"]}
        questions = json.loads((root / "evaluation" / "questions.json").read_text(encoding="utf-8"))["questions"]
        expected = json.loads((root / "evaluation" / "expected_sources.json").read_text(encoding="utf-8"))["expected_sources"]
        claims = json.loads((root / "metadata" / "claims_ledger.json").read_text(encoding="utf-8"))["claims"]
        answers = json.loads((root / "evaluation" / "answer_key.json").read_text(encoding="utf-8"))["answers"]
        expected_by_id = {item["question_id"]: item for item in expected}
        claims_by_id = {claim["claim_id"]: claim for claim in claims}
        answers_by_id = {answer["question_id"]: answer for answer in answers}

        self.assertEqual(len(questions), 20)
        self.assertEqual(set(expected_by_id), {item["question_id"] for item in questions})
        self.assertEqual(set(answers_by_id), {item["question_id"] for item in questions})
        for question in questions:
            question_id = question["question_id"]
            self.assertTrue(set(question["expected_source_ids"]) <= known_sources, question_id)
            self.assertEqual(question["expected_source_ids"], expected_by_id[question_id]["required_source_ids"])
            self.assertEqual(question["abstention_required"], answers_by_id[question_id]["abstention_required"])
            for claim_id in question["expected_claim_ids"]:
                self.assertIn(claim_id, claims_by_id, question_id)
                self.assertTrue(set(claims_by_id[claim_id]["supporting_source_ids"]) <= known_sources)
        for source in manifest["sources"]:
            parsed = parse_source(root, source)
            self.assertEqual(parsed.source["source_id"], source["source_id"])
            self.assertEqual(parsed.source_hash, source["sha256"])
            self.assertTrue(parsed.units, source["source_id"])

    def test_markdown_and_docx_parsers_preserve_section_locations(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            markdown = root / "guide.md"
            markdown.write_text("# Recovery\nRestart the worker.", encoding="utf-8")
            parsed_markdown = parse_source(root, {
                "source_id": "DOC-MD",
                "title": "Guide",
                "source_type": "document",
                "file_path": markdown.name,
            })
            self.assertEqual(parsed_markdown.units[1]["locator"]["section"], "Recovery")

            document = Document()
            document.add_heading("Operations", level=1)
            document.add_paragraph("Use a staged rollout.")
            document_path = root / "guide.docx"
            document.save(document_path)
            parsed_docx = parse_source(root, {
                "source_id": "DOCX-1",
                "title": "Operations guide",
                "source_type": "document",
                "file_path": document_path.name,
            })
            self.assertEqual(parsed_docx.units[1]["locator"]["section"], "Operations")

    def test_vtt_voice_tags_become_speakers(self):
        root, manifest = load_manifest(Path(__file__).resolve().parents[1] / "synthetic_data")
        entry = next(source for source in manifest["sources"] if source["source_id"] == "TRN-002")
        cue = next(unit for unit in parse_source(root, entry).units if unit["locator"]["cue_id"] == "cue-004")
        self.assertEqual(cue["metadata"]["speaker"], "Arjun Rao")
        self.assertTrue(cue["text"].startswith("Also, the HTTP status on these errors is 422"))


class ChunkingTests(unittest.TestCase):
    def test_chunks_are_deterministic_and_keep_locator(self):
        source = {
            "source_id": "DOC-X",
            "source_type": "document",
            "title": "Example",
            "source_hash": "abc",
            "incident_id": "INC-X",
        }
        unit = {"text": " ".join(f"word{i}" for i in range(20)), "locator": {"page": 4}}
        first = chunk_evidence_units([unit], source, chunk_size=8, overlap=2)
        second = chunk_evidence_units([unit], source, chunk_size=8, overlap=2)
        self.assertGreater(len(first), 1)
        self.assertEqual([item["chunk_id"] for item in first], [item["chunk_id"] for item in second])
        self.assertTrue(all(item["locator"] == {"page": 4} for item in first))
        self.assertTrue(all(item["source_id"] == "DOC-X" for item in first))
        self.assertTrue(all(len(item["content_hash"]) == 64 for item in first))

    def test_small_evidence_unit_stays_intact_with_locator(self):
        chunks = chunk_document_with_location(
            "A transcript cue is one searchable unit.",
            "TRN-1",
            SourceType.TRANSCRIPT,
            chunk_size=100,
            locator={"cue_id": "cue-1", "start": "00:00:01.000"},
        )
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0]["locator"]["cue_id"], "cue-1")

    def test_repeated_content_keeps_distinct_source_locators(self):
        source = {
            "source_id": "DOC-X",
            "source_type": "document",
            "source_hash": "abc",
        }
        units = [
            {"text": "The same evidence appears twice.", "locator": {"section": "First"}},
            {"text": "The same evidence appears twice.", "locator": {"section": "Second"}},
        ]
        chunks = chunk_evidence_units(units, source, chunk_size=100, overlap=0)
        self.assertEqual(chunks[0]["content_hash"], chunks[1]["content_hash"])
        self.assertNotEqual(chunks[0]["chunk_id"], chunks[1]["chunk_id"])
        self.assertEqual(
            [chunk["locator"]["section"] for chunk in chunks],
            ["First", "Second"],
        )


class RetrievalLogicTests(unittest.TestCase):
    @staticmethod
    def citation(source_id: str, chunk_id: str) -> Citation:
        return Citation(
            source_id=source_id,
            chunk_id=chunk_id,
            source_type=SourceType.DOCUMENT,
            title=source_id,
            location="Page 1",
            snippet="Evidence",
        )

    def test_rrf_deduplicates_and_rewards_cross_method_results(self):
        shared = self.citation("DOC-1", "chunk-shared")
        vector_only = self.citation("DOC-2", "chunk-vector")
        lexical_only = self.citation("REC-1", "chunk-lexical")
        fused = reciprocal_rank_fusion(
            [[shared, vector_only], [shared, lexical_only]],
            top_k=3,
        )
        self.assertEqual(fused[0].chunk_id, "chunk-shared")
        self.assertEqual(fused[0].metadata["retrieval_methods"], ["lexical", "vector"])
        self.assertEqual(len({citation.chunk_id for citation in fused}), 3)

    def test_vector_filter_builds_expected_metadata_predicates(self):
        actual = build_vector_filter({
            "incident_id": "INC-1",
            "version": "2.4.0",
            "owner": "payments-team",
            "source_types": [SourceType.DOCUMENT, SourceType.TRANSCRIPT],
            "time_range_start": "2026-10-01T00:00:00Z",
        })
        self.assertEqual(actual["$and"][0], {"incident_id": {"$eq": "INC-1"}})
        self.assertEqual(actual["$and"][1], {"version": {"$eq": "2.4.0"}})
        self.assertEqual(actual["$and"][2], {"owner": {"$eq": "payments-team"}})
        self.assertEqual(actual["$and"][3], {"source_type": {"$in": ["document", "transcript"]}})
        self.assertIn("timestamp", actual["$and"][4])

    def test_atlas_bm25_filters_use_search_compound_operators(self):
        actual = _atlas_search_filters({
            "incident_id": "INC-1",
            "owner": "payments-team",
            "source_types": [SourceType.DOCUMENT],
            "time_range_start": "2026-10-01T00:00:00Z",
        })
        self.assertEqual(actual[0], {"equals": {"path": "incident_id", "value": "INC-1"}})
        self.assertEqual(actual[1], {"equals": {"path": "owner", "value": "payments-team"}})
        self.assertEqual(actual[2], {"in": {"path": "source_type", "value": ["document"]}})
        self.assertEqual(
            actual[3]["range"]["gte"],
            __import__("datetime").datetime.fromisoformat(
                "2026-10-01T00:00:00+00:00"
            ),
        )

    def test_metrics_cover_precision_recall_mrr_and_ndcg(self):
        actual = retrieval_metrics(
            ["SRC-1", "SRC-1", "SRC-2"],
            ["SRC-2", "SRC-3"],
            top_k=3,
        )
        self.assertEqual(actual["retrieved_source_ids"], ["SRC-1", "SRC-2"])
        self.assertEqual(actual["recall_at_k"], 0.5)
        self.assertAlmostEqual(actual["precision_at_k"], 1 / 3)
        self.assertEqual(actual["reciprocal_rank"], 0.5)
        self.assertGreater(actual["ndcg_at_k"], 0.0)
        self.assertLess(actual["ndcg_at_k"], 1.0)


class QueryPlanningTests(unittest.TestCase):
    def test_rewrite_expands_common_abbreviations_without_changing_identifiers(self):
        self.assertEqual(
            rewrite_query("Why did PAY-401 occur in prod for svc?"),
            "Why did PAY-401 occur in production for service?",
        )

    def test_decomposition_keeps_original_question_and_extracts_second_hop(self):
        queries = decompose_query(
            "Why did payments fail after deployment, and what changed in the request schema?"
        )
        self.assertEqual(len(queries), 2)
        self.assertTrue(queries[0].startswith("Why did payments fail"))
        self.assertEqual(queries[1], "changed in the request schema?")


class GeminiServiceTests(unittest.TestCase):
    def test_embedding_adapters_use_retrieval_task_types_and_expected_dimensions(self):
        model = MagicMock()
        model.embed_content.return_value = SimpleNamespace(
            embeddings=[SimpleNamespace(values=[0.1] * 768)]
        )
        client = SimpleNamespace(models=model)
        with patch("app.services.embeddings.get_gemini_client", return_value=client):
            vectors = embed_texts(["document text"])
            query = embed_query("search question")
        self.assertEqual(len(vectors[0]), 768)
        self.assertEqual(len(query), 768)
        self.assertEqual(model.embed_content.call_args_list[0].kwargs["config"].task_type, "RETRIEVAL_DOCUMENT")
        self.assertEqual(model.embed_content.call_args_list[1].kwargs["config"].task_type, "RETRIEVAL_QUERY")

    def test_claim_verification_withholds_unsupported_claims(self):
        citation = RetrievalLogicTests.citation("DOC-1", "chunk-1")
        claim = SupportedClaim(
            claim="The release caused the error.",
            classification="inference",
            citation_chunk_ids=["chunk-1"],
        )
        with patch("app.reasoning.investigator.generate_json", return_value={
            "verifications": [{
                "claim_index": 0,
                "status": "insufficient",
                "explanation": "The cited text does not establish causality.",
            }],
        }):
            verified = verify_claims([claim], {"chunk-1": citation})
        self.assertEqual(verified[0].classification, "unknown")
        self.assertEqual(verified[0].citation_chunk_ids, [])
        self.assertEqual(verified[0].verification_status, "insufficient")

    def test_generation_uses_gemini_structured_json_mode(self):
        model = MagicMock()
        model.generate_content.return_value = SimpleNamespace(text='{"answer":"grounded"}')
        client = SimpleNamespace(models=model)
        with patch("app.services.gemini.get_gemini_client", return_value=client):
            result = generate_json("Use evidence only.", {"question": "What happened?"})
        self.assertEqual(result, {"answer": "grounded"})
        config = model.generate_content.call_args.kwargs["config"]
        self.assertEqual(config.response_mime_type, "application/json")
        self.assertEqual(config.system_instruction, "Use evidence only.")

    def test_answer_only_contains_claims_that_pass_independent_verification(self):
        citation = RetrievalLogicTests.citation("DOC-1", "chunk-1")
        generated = {
            "answer": "The payment service started returning PAY-401.",
            "likely_root_cause": None,
            "root_cause_chunk_ids": [],
            "is_abstention": False,
            "claims": [{
                "claim": "The payment service started returning PAY-401.",
                "classification": "confirmed",
                "chunk_ids": ["chunk-1"],
            }],
            "conflicting_chunk_ids": [],
            "recommended_actions": [],
            "uncertainties": [],
        }
        with patch("app.reasoning.investigator.generate_json", side_effect=[
            generated,
            {"verifications": [{
                "claim_index": 0,
                "status": "supported",
                "explanation": "The evidence directly states this.",
            }]},
        ]):
            report = investigate_query("What failed?", [citation])
        self.assertFalse(report.is_abstention)
        self.assertEqual(report.answer, generated["claims"][0]["claim"])
        self.assertEqual(report.claims[0].verification_status, "verified")
        self.assertEqual(report.supporting_evidence[0].chunk_id, "chunk-1")

    def test_unretrieved_generated_chunk_reference_causes_abstention(self):
        citation = RetrievalLogicTests.citation("DOC-1", "chunk-1")
        with patch("app.reasoning.investigator.generate_json", return_value={
            "answer": "Unsupported.",
            "likely_root_cause": None,
            "root_cause_chunk_ids": [],
            "is_abstention": False,
            "claims": [{
                "claim": "Unsupported claim.",
                "classification": "confirmed",
                "chunk_ids": ["invented-chunk"],
            }],
            "conflicting_chunk_ids": [],
            "recommended_actions": [],
            "uncertainties": [],
        }) as generate:
            report = investigate_query("What happened?", [citation])
        generate.assert_called_once()
        self.assertTrue(report.is_abstention)
        self.assertIn("verification", report.abstention_reason)
        self.assertEqual(report.supporting_evidence[0].chunk_id, "chunk-1")


class ApiContractTests(unittest.TestCase):
    def test_health_route_is_live(self):
        with TestClient(app) as client:
            response = client.get("/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "ok")

    def test_invalid_investigation_filter_is_rejected_before_services(self):
        with TestClient(app) as client:
            response = client.post("/api/investigate/", json={
                "query": "Find deployment details",
                "time_range_start": "2026-10-24T00:00:00",
            })
        self.assertEqual(response.status_code, 422)

    def test_upload_rejects_unsupported_format_before_database_access(self):
        with TestClient(app) as client:
            response = client.post(
                "/api/ingest/upload",
                files={"file": ("payload.exe", b"not an accepted evidence format")},
                data={"source_type": "document"},
            )
        self.assertEqual(response.status_code, 400)

    def test_index_timeout_setting_must_be_finite(self):
        for value in ("nan", "inf"):
            with self.assertRaises(ValueError):
                Settings(SEARCH_INDEX_TIMEOUT_SECONDS=value)

    def test_other_timeouts_are_not_reported_as_index_builds(self):
        with patch("app.api.routes_ingest.ingest_dataset", side_effect=TimeoutError("socket")):
            with TestClient(app) as client:
                response = client.post("/api/ingest/batch")
        self.assertEqual(response.status_code, 503)
        self.assertNotIn("still building", response.json()["detail"])

    def test_ingest_while_search_indexes_build_returns_503(self):
        with patch("app.api.routes_ingest.ingest_dataset", side_effect=SearchIndexNotReadyError("building")):
            with TestClient(app) as client:
                response = client.post("/api/ingest/batch")
        self.assertEqual(response.status_code, 503)
        self.assertIn("still building", response.json()["detail"])

    def test_upload_while_search_indexes_build_returns_503(self):
        with patch("app.api.routes_ingest.ingest_uploaded_file", side_effect=SearchIndexNotReadyError("building")):
            with TestClient(app) as client:
                response = client.post(
                    "/api/ingest/upload",
                    files={"file": ("notes.txt", b"STATUS\nMitigated.")},
                    data={"source_type": "transcript"},
                )
        self.assertEqual(response.status_code, 503)
        self.assertIn("still building", response.json()["detail"])


if __name__ == "__main__":
    unittest.main()
